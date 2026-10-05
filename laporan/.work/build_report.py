from pathlib import Path
from collections import Counter, defaultdict
import csv
import math
import statistics
import json
import hashlib
import shutil
from datetime import datetime

from PIL import Image, ImageDraw, ImageFont
import numpy as np
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE = Path('/home/parasite/Kuliah/JARNIR/kelompok4_L8')
OUT = BASE / 'laporan'
WORK = OUT / '.work'
ASSETS = WORK / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)


def read_csv(path):
    with path.open(newline='', encoding='utf-8') as stream:
        return list(csv.DictReader(stream))


passive = read_csv(BASE / 'data/passive_L8.csv')
active = read_csv(BASE / 'data/active_L8.csv')
points = read_csv(BASE / 'output/L8_titik.csv')
radios = read_csv(BASE / 'output/L8_ap.csv')
roam = read_csv(BASE / 'raw/roaming_L8.csv')
android = read_csv(BASE / 'data/android_L8_transkrip.csv')
locations = {r['titik_id']: r['keterangan'] for r in read_csv(BASE / 'data/titik_L8.csv')}
target = [r for r in passive if r['ssid'] == 'WiFi-UB.x' and r['band'] == '5']
names = {
    'L8-001': 'Koridor kanan bawah',
    'L8-002': 'Koridor kanan atas',
    'L8-003': 'Koridor atas tengah',
    'L8-004': 'Koridor atas kiri',
    'L8-005': 'Koridor tengah kiri',
    'L8-006': 'Koridor bawah kiri dekat tangga',
    'L8-007': 'Koridor bawah tengah',
    'L8-008': 'Lab RPL',
    'L8-009': 'Depan lift sisi kiri',
    'L8-010': 'Depan lift sisi kanan',
    'L8-011': 'Toilet',
    'L8-012': 'Koridor di bawah AP-L8-02',
}


def n(value, digits=1):
    if value is None or value == '':
        return 'Tidak terdeteksi'
    return f'{float(value):.{digits}f}'.replace('.', ',')


def linear(values):
    return 10 * math.log10(sum(10 ** (v / 10) for v in values) / len(values))


# Verify the principal metrics independently from the raw observations.
grouped = defaultdict(list)
for row in target:
    grouped[row['titik_id'], row['bssid'].lower()].append(float(row['rssi_dbm']))
derived = {}
for tid in sorted(names):
    signal = sorted([linear(v) for (t, b), v in grouped.items() if t == tid], reverse=True)
    derived[tid] = signal
    stored = next(p for p in points if p['titik_id'] == tid)
    assert round(signal[0], 1) == float(stored['rssi1'])
    if len(signal) > 1:
        assert round(signal[1], 1) == float(stored['rssi2'])
    else:
        assert stored['rssi2'] == ''
assert len(passive) == 428 and len(target) == 75 and len(active) == 6
assert len(roam) == 145 and sum(r['pindah'] == 'YA' for r in roam) == 3
ssid_by_bssid = defaultdict(set)
for row in passive:
    ssid_by_bssid[row['bssid'].lower()].add(row['ssid'])
assert all(ssid_by_bssid[r['bssid_assoc']] == {'eduroam'} for r in active)
assert all(ssid_by_bssid[r['bssid']] == {'eduroam'} for r in roam if r['bssid'])

# Recreate all report charts from stored observations and retain the floor plan's aspect ratio.
font = '/usr/share/fonts/noto/NotoSans-Regular.ttf'
bold = '/usr/share/fonts/noto/NotoSans-Bold.ttf'
F = lambda size, b=False: ImageFont.truetype(bold if b else font, size)
plan = Image.open(BASE / 'denah/sketsa_L8.jpg').convert('RGB')
pw, ph = plan.size
map_width = 1420
map_height = round(ph * map_width / pw)
origin_x, origin_y = 80, 105
xy_source = np.array([[float(r['x_px']) * pw / 920, float(r['y_px']) * ph / 920] for r in points])
xy_plot = xy_source * map_width / pw + np.array([origin_x, origin_y])
values = np.array([float(r['rssi1']) for r in points])


def triangulate(coords):
    # Bowyer-Watson Delaunay triangulation for the small set of survey points.
    center = coords.mean(axis=0)
    span = np.ptp(coords, axis=0).max() * 25
    nodes = np.vstack([coords, center + [-2 * span, -span], center + [0, 2 * span], center + [2 * span, -span]])
    size = len(coords)
    triangles = [(size, size + 1, size + 2)]
    def within(t, point):
        a, b, c = nodes[list(t)]
        lhs = 2 * np.array([b - a, c - a])
        rhs = np.array([np.dot(b, b) - np.dot(a, a), np.dot(c, c) - np.dot(a, a)])
        try:
            mid = np.linalg.solve(lhs, rhs)
        except np.linalg.LinAlgError:
            return False
        return np.dot(point - mid, point - mid) <= np.dot(a - mid, a - mid) + 1e-5
    for idx, point in enumerate(coords):
        removed = [t for t in triangles if within(t, point)]
        edges = Counter(tuple(sorted((t[i], t[(i + 1) % 3]))) for t in removed for i in range(3))
        triangles = [t for t in triangles if t not in removed]
        triangles.extend((edge[0], edge[1], idx) for edge, count in edges.items() if count == 1)
    return [t for t in triangles if all(v < size for v in t)]


grid_width = 900
grid_height = round(ph * grid_width / pw)
gx, gy = np.meshgrid(np.linspace(0, pw, grid_width), np.linspace(0, ph, grid_height))
grid = np.full(gx.shape, np.nan)
tris = triangulate(xy_source)
assert len(tris) > 0
for t in tris:
    a, b, c = xy_source[list(t)]
    det = (b[1] - c[1]) * (a[0] - c[0]) + (c[0] - b[0]) * (a[1] - c[1])
    if abs(det) < 1e-9:
        continue
    wa = ((b[1] - c[1]) * (gx - c[0]) + (c[0] - b[0]) * (gy - c[1])) / det
    wb = ((c[1] - a[1]) * (gx - c[0]) + (a[0] - c[0]) * (gy - c[1])) / det
    wc = 1 - wa - wb
    mask = (wa >= -1e-7) & (wb >= -1e-7) & (wc >= -1e-7)
    grid[mask] = wa[mask] * values[t[0]] + wb[mask] * values[t[1]] + wc[mask] * values[t[2]]

color_values = np.array([-85, -80, -70, -67, -55, -40])
colors = np.array([[165, 0, 38], [215, 48, 39], [253, 174, 97], [255, 255, 191], [102, 189, 99], [0, 104, 55]])


def rssi_color(v):
    return tuple(int(np.interp(v, color_values, colors[:, i])) for i in range(3))


def new_map(title, heat=False):
    im = Image.new('RGB', (1580, map_height + 265), 'white')
    d = ImageDraw.Draw(im)
    d.text((origin_x, 18), title, fill='black', font=F(29, True))
    d.text((origin_x, 60), 'WiFi-UB.x 5 GHz | denah skematis | nilai pada titik berasal dari scan pasif', fill='#555555', font=F(22))
    background = plan.resize((map_width, map_height), Image.Resampling.LANCZOS).convert('RGBA')
    if heat:
        valid = np.isfinite(grid)
        rgba = np.zeros((grid_height, grid_width, 4), dtype=np.uint8)
        for channel in range(3):
            rgba[:, :, channel] = np.interp(np.nan_to_num(grid, nan=-85), color_values, colors[:, channel]).astype(np.uint8)
        rgba[:, :, 3] = np.where(valid, 110, 0).astype(np.uint8)
        overlay = Image.fromarray(rgba, 'RGBA').resize((map_width, map_height), Image.Resampling.BILINEAR)
        background = Image.alpha_composite(background, overlay)
    im.paste(background.convert('RGB'), (origin_x, origin_y))
    return im, ImageDraw.Draw(im)


heatmap, hd = new_map('Heatmap RSSI AP terkuat pada lantai 8', True)
for (x, y), row in zip(xy_plot, points):
    hd.ellipse((x - 10, y - 10, x + 10, y + 10), fill=rssi_color(float(row['rssi1'])), outline='black', width=2)
    text = n(row['rssi1'])
    label_box = hd.textbbox((x + 15, y - 16), text, font=F(21, True))
    hd.rectangle((label_box[0] - 3, label_box[1] - 3, label_box[2] + 3, label_box[3] + 3), fill='white')
    hd.text((x + 15, y - 16), text, fill='black', font=F(21, True))
legend_y = origin_y + map_height + 28
for i in range(850):
    val = -85 + i / 849 * 45
    hd.line([(origin_x + i, legend_y), (origin_x + i, legend_y + 22)], fill=rssi_color(val), width=1)
for val in [-85, -80, -70, -67, -55, -40]:
    x = origin_x + (val + 85) / 45 * 849
    hd.text((x - 19, legend_y + 29), str(val), fill='black', font=F(19))
hd.text((origin_x + 890, legend_y - 1), 'RSSI (dBm)', fill='black', font=F(22))
hd.text((origin_x, legend_y + 66), 'Interpolasi linear hanya di dalam batas titik ukur; area tanpa warna belum dipetakan.', fill='#555555', font=F(21))
heatmap_file = BASE / 'output/L8_heatmap_laporan.png'
heatmap.save(heatmap_file)

status_map, sd = new_map('Peta status titik ukur dan observasi AP')
status_colors = {'BAIK': '#1a9850', 'PERHATIAN': '#fee08b', 'BURUK': '#f46d43', 'KRITIS': '#a50026'}
for (x, y), row in zip(xy_plot, points):
    sd.ellipse((x - 14, y - 14, x + 14, y + 14), fill=status_colors[row['status']], outline='black', width=2)
physical = read_csv(BASE / 'data/ap_fisik_L8.csv')
for row in physical:
    x = origin_x + float(row['x_px']) / 920 * map_width
    y = origin_y + float(row['y_px']) / 920 * map_height
    sd.rectangle((x - 13, y - 13, x + 13, y + 13), outline='#2266a8', width=4)
for row in radios:
    x = origin_x + float(row['x_est']) / 920 * map_width
    y = origin_y + float(row['y_est']) / 920 * map_height
    sd.polygon([(x, y - 16), (x - 15, y + 12), (x + 15, y + 12)], fill='#2266a8')
    sd.text((x + 18, y - 17), 'ch' + row['channel'], fill='#2266a8', font=F(21, True))
legend_y = origin_y + map_height + 26
for i, status in enumerate(['BAIK', 'PERHATIAN', 'BURUK', 'KRITIS']):
    x = origin_x + i * 335
    sd.ellipse((x, legend_y, x + 22, legend_y + 22), fill=status_colors[status], outline='black', width=2)
    count = sum(r['status'] == status for r in points)
    sd.text((x + 35, legend_y - 5), f'{status} ({count})', fill='black', font=F(21))
sd.text((origin_x, legend_y + 47), 'Kotak biru: AP fisik   |   Segitiga biru: posisi radio perkiraan   |   Koordinat bersifat skematis', fill='#555555', font=F(21))
status_file = BASE / 'output/L8_status_laporan.png'
status_map.save(status_file)

channel_plot = Image.new('RGB', (1500, 620), 'white')
cd = ImageDraw.Draw(channel_plot)
cd.text((75, 20), 'Kanal radio WiFi-UB.x yang terdeteksi', fill='black', font=F(32, True))
cd.text((75, 69), 'Tiga radio target | masing-masing 20 MHz | maksimum RSSI pada titik pengamatan', fill='#555555', font=F(22))
baseline = 475
for index, row in enumerate(radios):
    x = 240 + index * 430
    cd.rectangle((x - 110, 195, x + 110, baseline), fill=['#287ba8', '#489a7c', '#7855a1'][index])
    cd.text((x - 38, 242), '1 radio', fill='white', font=F(25, True))
    cd.text((x - 82, 305), n(row['rssi_maks']) + ' dBm', fill='white', font=F(24))
    cd.text((x - 68, baseline + 20), 'Kanal ' + row['channel'], fill='black', font=F(25, True))
    cd.text((x - 95, baseline + 61), row['radio'][-8:].upper(), fill='#555555', font=F(22))
cd.text((75, 575), 'Jumlah radio target tidak sama dengan jumlah seluruh BSSID atau jumlah AP resmi gedung.', fill='#555555', font=F(21))
channel_file = BASE / 'output/L8_kanal_laporan.png'
channel_plot.save(channel_file)

# Preserve both supplied photo files as original evidence, without modifying them.
photo_one_source = Path('/home/parasite/Downloads/WhatsApp Image 2026-10-05 at 10.09.38.jpeg')
photo_other_source = Path('/tmp/codex-clipboard-32fe707a-883d-4fde-a387-7f4e157cf676.png')
photo_one = BASE / 'denah/AP-L8-01.jpeg'
photo_other = BASE / 'denah/AP-L8-02_dan_AP-L8-03.png'
for source, dest in [(photo_one_source, photo_one), (photo_other_source, photo_other)]:
    if not dest.exists():
        shutil.copy2(source, dest)
    assert hashlib.sha256(source.read_bytes()).digest() == hashlib.sha256(dest.read_bytes()).digest()

# A scientific time-series plot of the stored roaming observations.
font = '/usr/share/fonts/noto/NotoSans-Regular.ttf'
bold = '/usr/share/fonts/noto/NotoSans-Bold.ttf'
canvas = Image.new('RGB', (1600, 680), 'white')
draw = ImageDraw.Draw(canvas)
F = lambda size, b=False: ImageFont.truetype(bold if b else font, size)
left, top, right, bottom = 130, 125, 1525, 495
start = datetime.strptime(roam[0]['waktu'], '%H:%M:%S')
seconds = lambda text: (datetime.strptime(text, '%H:%M:%S') - start).total_seconds()
palette = {
    '00:e6:3a:10:a9:5c': '#176a9a',
    '00:e6:3a:10:ea:5c': '#b4770c',
    '00:e6:3a:10:eb:0c': '#7d48a3',
}
X = lambda sec: left + sec / 146 * (right - left)
Y = lambda dbm: bottom - (dbm + 90) / 55 * (bottom - top)
draw.text((left, 18), 'RSSI koneksi selama pencatatan roaming', fill='black', font=F(32, True))
draw.text((left, 63), 'BSSID terasosiasi cocok dengan eduroam pada hasil scan pasif', fill='#555555', font=F(23))
for val in [-90, -80, -70, -60, -50, -40]:
    yy = Y(val)
    draw.line([(left, yy), (right, yy)], fill='#d9dfe4', width=1)
    draw.text((left - 70, yy - 16), str(val), fill='black', font=F(23))
draw.text((10, 95), 'dBm', fill='black', font=F(23))
for sec in [0, 30, 60, 90, 120, 146]:
    xx = X(sec)
    draw.line([(xx, top), (xx, bottom)], fill='#edf0f2', width=1)
    label = (start.timestamp() + sec)
    stamp = datetime.fromtimestamp(label).strftime('%H:%M:%S')
    draw.text((xx - 49, bottom + 20), stamp, fill='black', font=F(21))
prev = None
for row in roam:
    if not row['bssid']:
        xx = X(seconds(row['waktu']))
        draw.line([(xx, top), (xx, bottom)], fill='#a0a0a0', width=3)
        prev = None
        continue
    xy = X(seconds(row['waktu'])), Y(float(row['rssi_dbm']))
    if prev and prev[0] == row['bssid']:
        draw.line([prev[1], xy], fill=palette[row['bssid']], width=4)
    draw.ellipse((xy[0] - 3, xy[1] - 3, xy[0] + 3, xy[1] + 3), fill=palette[row['bssid']])
    if row['pindah'] == 'YA':
        draw.line([(xy[0], top), (xy[0], bottom)], fill='#555555', width=1)
    prev = row['bssid'], xy
for i, (bssid, color) in enumerate(palette.items()):
    xx = left + i * 450
    draw.line([(xx, 585), (xx + 45, 585)], fill=color, width=6)
    ch = {40: 'A9:5C', 48: 'EA:5C', 161: 'EB:0C'}
    channel = {'a9:5c': 40, 'ea:5c': 48, 'eb:0c': 161}[bssid[-5:]]
    draw.text((xx + 58, 565), f'{bssid[-5:].upper()}   kanal {channel}', fill='black', font=F(23))
draw.text((left, 627), 'Garis abu-abu: BSSID kosong. Garis vertikal gelap: perubahan BSSID yang tercatat.', fill='#555555', font=F(20))
roam_plot = BASE / 'output/L8_roaming_laporan.png'
canvas.save(roam_plot)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

for style_name in ['Normal', 'Body Text']:
    st = doc.styles[style_name]
    st.font.name = 'Noto Serif'
    st.font.size = Pt(11)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.13
    st.paragraph_format.widow_control = True
for style_name, size in [('Title', 23), ('Subtitle', 12), ('Heading 1', 15), ('Heading 2', 12)]:
    st = doc.styles[style_name]
    st.font.name = 'Noto Sans'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = style_name.startswith('Heading')
    st.paragraph_format.space_before = Pt(10 if style_name.startswith('Heading') else 0)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.line_spacing = 1.05
    st.font.underline = False
    ppr = st.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in list(ppr.findall(qn('w:pBdr'))):
            ppr.remove(border)
doc.styles['Caption'].font.name = 'Noto Sans'
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.color.rgb = RGBColor(70, 70, 70)
doc.styles['Caption'].paragraph_format.space_after = Pt(9)

header = section.header.paragraphs[0]
header.text = 'JARINGAN NIRKABEL   |   KELOMPOK 4   |   LANTAI 8'
header.style = 'Caption'
header.paragraph_format.space_after = Pt(0)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('Halaman ')
run.font.name = 'Noto Sans'
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def p(text, bold_lead=None):
    par = doc.add_paragraph(style='Normal')
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        par.add_run(bold_lead).bold = True
        par.add_run(text[len(bold_lead):])
    else:
        par.add_run(text)
    return par


def h(text, level=1):
    return doc.add_heading(text, level)


def page(title):
    heading = h(title)
    heading.paragraph_format.page_break_before = True


def table(headers, rows, widths, font_size=9.6, left_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    properties = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + name)
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:color'), 'D9D9D9')
        borders.append(border)
    properties.append(borders)
    for row_i, values in enumerate([headers] + rows):
        cells = t.rows[0].cells if row_i == 0 else t.add_row().cells
        tr_pr = cells[0]._tc.getparent().get_or_add_trPr()
        tr_pr.append(OxmlElement('w:cantSplit'))
        if row_i == 0:
            tr_pr.append(OxmlElement('w:tblHeader'))
        for col_i, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for name, val in [('top', 70), ('bottom', 70), ('left', 85), ('right', 85)]:
                node = OxmlElement('w:' + name)
                node.set(qn('w:w'), str(val))
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            tcpr.append(margins)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'E7EDF2' if row_i == 0 else ('F5F7F9' if row_i % 2 == 0 else 'FFFFFF'))
            tcpr.append(shading)
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_i in left_cols else WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.line_spacing = 1.05
            par.paragraph_format.keep_with_next = row_i == 0
            run = par.add_run(str(value))
            run.font.name = 'Noto Sans'
            run.font.size = Pt(font_size)
            run.font.bold = row_i == 0
            run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def figure(path, caption, width=16.3, height=None):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.keep_with_next = True
    if height:
        inline = par.add_run().add_picture(str(path), height=Cm(height))
    else:
        inline = par.add_run().add_picture(str(path), width=Cm(width))
    inline._inline.docPr.set('descr', caption)
    cap = doc.add_paragraph(caption, style='Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


# Page 1
title = doc.add_paragraph('Laporan Pemetaan Kualitas Sinyal WiFi 5 GHz', style='Title')
doc.add_paragraph('Lantai 8 Gedung F FILKOM Universitas Brawijaya', style='Subtitle')
doc.add_paragraph('Mata Kuliah Jaringan Nirkabel   |   Kelompok 4', style='Subtitle')
h('1 Identitas dan ringkasan')
table(['Komponen', 'Keterangan'], [
    ['Kelompok', 'Kelompok 4'],
    ['Nama dan NIM anggota', '[Isi nama dan NIM anggota kelompok]'],
    ['Lokasi', 'Lantai 8 Gedung F FILKOM UB'],
    ['Tanggal pengukuran', '28 September 2026'],
    ['Waktu pencatatan', 'Pasif 15.56.53-16.04.59; aktif 16.13.23-16.18.16; roaming 16.21.04-16.23.30'],
    ['Fokus survei', 'Cakupan SSID WiFi-UB.x pada band 5 GHz'],
], [4.0, 12.6], 10, (0, 1))
p('Survei ini memetakan kualitas sinyal Wi-Fi pada 12 titik di lantai 8 untuk mengenali area lemah, menghubungkan hasil pengukuran dengan penempatan access point (AP), dan menyusun rekomendasi perbaikan. Pengamatan mencakup koridor, Lab RPL, area lift, toilet, serta titik di bawah AP koridor kanan.')
p('Sebanyak 11 dari 12 titik (91,7%) memiliki RSSI AP terkuat WiFi-UB.x sekurang-kurangnya -67 dBm. Status gabungan terdiri atas 10 titik BAIK, satu PERHATIAN, dan satu BURUK. Toilet pada L8-011 menjadi area lemah utama dengan RSSI -73,0 dBm dan tanpa AP cadangan WiFi-UB.x yang terdeteksi. L8-012 memiliki sinyal sangat kuat (-41,6 dBm), tetapi juga belum menunjukkan AP cadangan untuk SSID tersebut.')
p('Tiga radio WiFi-UB.x teridentifikasi pada kanal 40, 48, dan 161 dengan lebar kanal 20 MHz. Prioritas perbaikan adalah memperluas cakupan menuju toilet dan memeriksa kesinambungan cakupan antar-AP. Persentase tersebut menggambarkan titik yang diukur, bukan persentase luas seluruh lantai.')
p('Uji aktif dan roaming dicatat sebagai bukti tambahan pada koneksi eduroam berdasarkan pencocokan BSSID. Karena itu, hasilnya dipisahkan dari kesimpulan cakupan WiFi-UB.x. Log roaming menunjukkan tiga perubahan BSSID, tetapi kehilangan paket saat perpindahan tidak dapat dihitung karena log ping serentak tidak tersimpan.')

# Page 2
page('2 Metodologi survei')
h('2 1 Perangkat dan sumber data', 2)
p('Survei pasif dilakukan menggunakan satu laptop yang diberi kode laptopA, dengan pembacaan Wi-Fi melalui toolchain Linux dan iw. Tangkapan layar WiFiAnalyzer pada Android digunakan sebagai pembanding. Model laptop, model adaptor Wi-Fi, dan model ponsel tidak dicantumkan dalam berkas pengukuran, sehingga tidak digunakan sebagai dasar penjelasan perbedaan perangkat.')
table(['Sumber', 'Jumlah', 'Penggunaan'], [
    ['Scan pasif', '428 baris', '276 baris 5 GHz dan 152 baris 2,4 GHz'],
    ['WiFi-UB.x pada 5 GHz', '75 baris', 'Perhitungan RSSI, AP terkuat, AP cadangan, kanal'],
    ['Titik pengamatan', '12 titik', 'Tiga siklus scan pada setiap titik'],
    ['Uji aktif', '6 titik', 'RSSI koneksi, RTT, jitter, dan packet loss'],
    ['Log roaming', '145 sampel', 'BSSID, RSSI koneksi, dan perubahan asosiasi'],
], [4.5, 2.6, 9.5], 10, (0, 2))
h('2 2 Prosedur dan pengolahan', 2)
p('Setiap titik menggunakan nomor yang sama pada sketsa, data pasif, dan data aktif. Scan pasif dilakukan tiga kali; jumlah AP yang terlihat dapat berbeda antar-scan. Data 2,4 GHz tetap disimpan sebagai data mentah, tetapi tidak disertakan dalam analisis utama. Pengujian aktif memakai ICMP menuju filkom.ub.ac.id dengan 50 paket per titik. Throughput iperf3 tidak diukur.')
p('RSSI tiap BSSID dihitung dengan rata-rata daya linear, kemudian dikembalikan ke satuan dBm. Jika P_i = 10^(RSSI_i/10) mW, maka RSSI rata-rata = 10 log10((P_1 + ... + P_n)/n). AP utama dan cadangan dipilih dari BSSID WiFi-UB.x yang memiliki rata-rata daya tertinggi dan kedua tertinggi.')
table(['Parameter', 'Kriteria yang digunakan'], [
    ['RSSI AP utama', 'BAIK >= -67 dBm; PERHATIAN -70 sampai < -67; BURUK -80 sampai < -70; KRITIS < -80 atau tidak terdeteksi'],
    ['AP cadangan', 'Tidak terdeteksi atau < -75 dBm menjadi PERHATIAN'],
    ['CCI', '> 2 radio lain pada kanal yang overlap menjadi PERHATIAN'],
    ['Utilisasi kanal', '> 50% menjadi PERHATIAN'],
    ['Uji aktif', 'RTT rata-rata > 50 ms menjadi PERHATIAN; loss > 2% menjadi BURUK'],
], [3.7, 12.9], 9.4, (0, 1))
p('Status akhir mengikuti kondisi terburuk dari metrik yang tersedia, sesuai analyze.py. Noise floor kosong pada data ini, sehingga SNR tidak dihitung. Pengelompokan beberapa SSID menjadi radio yang sama menggunakan heuristik kanal, kemiripan BSSID, dan RSSI; hasilnya bukan inventaris resmi perangkat.')

# Page 3
page('3 Denah dan titik pengamatan')
figure(BASE / 'denah/sketsa_L8.jpg', 'Gambar 1 Denah lantai 8 dengan nomor titik pengamatan dan tiga AP fisik', width=16.3)
table(['Titik', 'Lokasi pengamatan'], [[tid, names[tid]] for tid in sorted(names)], [3.0, 13.6], 9.8, (1,))
p('Denah digunakan untuk menghubungkan temuan dengan lokasi pengamatan. Titik koordinat pada keluaran analisis diperoleh dari penempatan perkiraan pada gambar. Skala meter, jarak antar-AP, dan luas area lemah belum diverifikasi melalui dimensi lapangan; denah dan peta keluaran dibaca sebagai representasi skematis.')

# Page 4
page('4 Hasil pengukuran pada titik survei')
table(['Metrik', 'Hasil'], [
    ['Titik dengan RSSI >= -67 dBm', '11/12 atau 91,7%'],
    ['Status gabungan BAIK', '10/12 atau 83,3%'],
    ['Status PERHATIAN dan BURUK', 'Masing-masing 1/12 atau 8,3%'],
    ['Blank spot pada titik ukur', '0/12 atau 0,0%'],
    ['Median dan minimum RSSI AP utama', '-53,6 dBm dan -73,0 dBm'],
    ['AP cadangan >= -75 dBm', '10/12 atau 83,3%'],
    ['Titik dengan CCI > 2 radio', '0/12 atau 0,0%'],
], [10.0, 6.6], 10, (0,))
table(['Titik', 'RSSI utama\n(dBm)', 'RSSI kedua\n(dBm)', 'Kanal\nutama', 'Utilisasi\n(%)', 'Status'], [
    [r['titik_id'], n(r['rssi1']), n(r['rssi2']) if r['rssi2'] else '-', r['channel1'], n(r['util1']), r['status']]
    for r in points
], [2.6, 2.9, 2.9, 2.1, 2.6, 3.5], 9.1)
p('Tanda - berarti tidak ada BSSID WiFi-UB.x kedua yang terdeteksi pada data 5 GHz di titik tersebut. RSSI utama menunjukkan sinyal terkuat hasil scan pasif, bukan selalu AP yang sedang dipakai laptop. Utilisasi merupakan rata-rata nilai BSS Load yang tersedia untuk AP utama, bukan pengukuran menyeluruh seluruh trafik kanal.')
p('L8-012 menunjukkan perbedaan antara sinyal kuat dan status keseluruhan: RSSI memenuhi ambang BAIK, tetapi ketiadaan AP cadangan menjadikan statusnya PERHATIAN. Oleh sebab itu, 91,7% titik dengan RSSI baik tidak sama dengan 83,3% titik berstatus BAIK.')

# Page 5
page('4 1 Heatmap RSSI WiFi UB x')
figure(heatmap_file, 'Gambar 2 Heatmap RSSI AP terkuat WiFi-UB.x pada 5 GHz dengan interpolasi linear', width=16.2)
p('Heatmap menunjukkan sebaran nilai RSSI yang diinterpolasi dari 12 titik pengamatan. Hasil paling kuat tercatat di L8-012 pada koridor di bawah AP-L8-02, sedangkan hasil paling lemah tercatat di toilet L8-011. Koridor dan Lab RPL yang diukur umumnya memiliki RSSI di atas ambang -67 dBm.')
p('Perubahan warna di antara titik merupakan hasil interpolasi, bukan scan tambahan. Ruang yang tidak memiliki titik ukur tidak dapat dinyatakan berkualitas baik hanya berdasarkan warna heatmap. Interpolasi juga tidak menghitung redaman setiap dinding, pintu, atau partisi secara terpisah.')
p('Peta untuk laporan mempertahankan proporsi denah asli. Koordinat perkiraan pada berkas titik ditransformasikan dari kanvas analisis ke ukuran denah tersebut. Karena ukuran lapangan belum dikalibrasi, warna heatmap tidak digunakan untuk menyimpulkan luas cakupan dalam meter persegi atau jarak fisik yang presisi. Angka RSSI dan status pada titik pengamatan tetap menjadi dasar utama interpretasi.')

# Page 6
page('4 2 Peta status dan pemetaan AP')
figure(status_file, 'Gambar 3 Peta status titik dan perbandingan posisi AP perkiraan dengan observasi fisik', width=16.2)
table(['AP fisik', 'Kandidat radio\nWiFi-UB.x', 'Kanal', 'Dasar pencocokan'], [
    ['AP-L8-01', '00:e6:3a:90:a9:5c', '40', 'Posisi perkiraan mendekati AP pada blok tengah bawah'],
    ['AP-L8-02', '00:e6:3a:90:eb:0c', '161', 'RSSI -41,6 dBm di L8-012 tepat pada area AP koridor kanan'],
    ['AP-L8-03', '00:e6:3a:90:ea:5c', '48', 'Posisi perkiraan berada di sisi kiri blok tengah atas'],
], [2.6, 5.0, 1.6, 7.4], 9.0, (1, 3))
p('Pencocokan di atas merupakan kandidat berdasarkan sketsa dan pola RSSI, bukan pembacaan label MAC perangkat. Posisi radio pada analyze.py ditaksir dengan centroid berbobot dari tiga titik terkuat. Ketiga radio memiliki kategori keyakinan tinggi karena RSSI maksimum mencapai sekurang-kurangnya -55 dBm; kategori tersebut tidak membuktikan koordinat, identitas perangkat, atau lantai asal secara pasti.')
p('Ketiga AP memiliki dokumentasi foto pada lampiran. Tulisan Ruckus terlihat pada foto AP-L8-01. Model AP, MAC pada label fisik, serta konfigurasi daya pancar belum terverifikasi. Marker AP fisik pada peta keluaran juga mengikuti koordinat perkiraan.')

# Page 7
page('4 3 Kanal dan diagnosis area lemah')
figure(channel_file, 'Gambar 4 Distribusi kanal tiga radio WiFi-UB.x pada 5 GHz', width=16.2)
p('Tiga radio target WiFi-UB.x menggunakan kanal 40, 48, dan 161 dengan lebar kanal 20 MHz. Pada lebar tersebut, kanal ketiga radio target tidak overlap. Nilai CCI pada tabel titik adalah 0 di seluruh titik, dan tidak ada titik yang melewati ambang CCI > 2 radio pada definisi analisis ini. Temuan ini tidak membuktikan ketiadaan seluruh bentuk interferensi atau kepadatan jaringan pada jam lain.')
p('Utilisasi kanal AP utama berkisar 4,2-27,7%, di bawah ambang 50%. Namun, pengukuran hanya dilakukan pada satu sesi sore; kesimpulan mengenai kepadatan jam sibuk memerlukan pengamatan pada waktu lain.')
h('Toilet pada L8 011', 2)
p('L8-011 memiliki RSSI utama -73,0 dBm pada kanal 48 dan tidak menunjukkan AP cadangan WiFi-UB.x. Hasil ini mendukung diagnosis cakupan lemah pada titik toilet. Dinding dan posisi AP yang terpisah dari area toilet menjadi penjelasan yang masuk akal dari denah, tetapi besarnya redaman dan jarak fisik belum diukur. Karena CCI dan utilisasi tidak melampaui ambang pada sesi ini, bukti lebih mendukung masalah cakupan daripada kanal padat.')
h('Koridor di bawah AP L8 02 pada L8 012', 2)
p('L8-012 memiliki RSSI utama -41,6 dBm dan utilisasi 15,1%, tetapi hanya satu BSSID WiFi-UB.x yang terlihat pada 5 GHz. Sinyal utama cukup kuat; masalah yang ditandai adalah ketiadaan AP cadangan dalam sampel. Status PERHATIAN menyatakan potensi risiko kesinambungan cakupan, bukan bukti bahwa handover gagal di lokasi ini.')

# Page 8
page('4 4 Hasil uji aktif')
p('Uji aktif dilakukan pada L8-001, L8-007, L8-008, L8-009, L8-011, dan L8-012. BSSID koneksi yang dicatat memakai pola 00:e6:3a:10:... dan cocok dengan SSID eduroam pada data pasif. Hasil berikut menggambarkan koneksi eduroam yang terasosiasi saat pengujian, sehingga tidak dianggap sebagai pengujian layanan WiFi-UB.x secara langsung.')
table(['Titik', 'Kanal\nkoneksi', 'RSSI\n(dBm)', 'RTT rata\n(ms)', 'RTT maks\n(ms)', 'Loss\n(%)'], [
    [r['titik_id'], {5200: 40, 5240: 48, 5805: 161}[int(r['freq_mhz'])], n(r['rssi_dbm'], 0), n(r['rtt_avg_ms'], 2), n(r['rtt_max_ms'], 1), n(r['loss_pct'], 1)]
    for r in sorted(active, key=lambda r: r['titik_id'])
], [2.8, 2.4, 2.5, 3.0, 3.0, 2.9], 9.5)
p('RTT rata-rata berada pada rentang 4,03-9,82 ms. Lima titik tidak mencatat kehilangan paket dalam 50 paket yang dikirim. Toilet L8-011 mencatat satu paket tidak diterima dari 50 paket, setara 2,0%, dengan RSSI koneksi -79 dBm. Catatan ping mentah pada titik tersebut mengonfirmasi 50 paket dikirim dan 49 diterima.')
p('L8-012 memiliki RTT maksimum 119,0 ms meskipun RTT rata-ratanya 9,82 ms dan packet loss 0,0%. Saat uji aktif, laptop terhubung ke radio kanal 48 dengan RSSI -75 dBm, sedangkan scan pasif WiFi-UB.x menunjukkan radio kanal 161 yang jauh lebih kuat. Perbedaan ini menunjukkan bahwa sinyal AP terkuat yang tersedia tidak selalu sama dengan AP yang dipilih klien; SSID dan waktu pengukuran juga berbeda.')
p('Loss 2,0% pada toilet tepat berada di batas, bukan melebihi ambang > 2% pada modul. Status BURUK L8-011 berasal dari RSSI pasif yang rendah. Hasil aktif memberi bukti tambahan bahwa koneksi di lokasi tersebut patut diperhatikan, tetapi satu pengujian singkat tidak cukup untuk menentukan frekuensi kehilangan paket sehari-hari.')
p('Kolom throughput downlink dan uplink kosong karena iperf3 tidak dijalankan. Nilai tx_rate_mbps adalah laju link yang dilaporkan adaptor, bukan throughput aplikasi. Laporan ini tidak menyimpulkan kecepatan unduh atau unggah dari nilai tersebut.')

# Page 9
page('4 5 Hasil pencatatan roaming')
p('Log roaming berisi 145 sampel dari 16.21.04 sampai 16.23.30, dengan selang waktu awal-akhir 146 detik. BSSID yang terasosiasi cocok dengan eduroam pada hasil scan pasif. Tiga perubahan BSSID tercatat selama lintasan; waktu peristiwa merupakan waktu sampel, bukan pengukuran durasi handover.')
figure(roam_plot, 'Gambar 5 Perubahan RSSI koneksi dan BSSID selama pencatatan roaming', width=16.2)
table(['Waktu\nperubahan', 'BSSID lama ke baru\nakhiran MAC', 'Kanal\nlama ke baru', 'RSSI sebelum\n(dBm)', 'RSSI sesudah\n(dBm)'], [
    ['16.21.09', 'A9:5C ke EA:5C', '40 ke 48', '-56', '-50'],
    ['16.21.46', 'EA:5C ke EB:0C', '48 ke 161', '-71', '-66'],
    ['16.22.56', 'EB:0C ke A9:5C', '161 ke 40', '-78', '-42'],
], [2.6, 4.6, 2.8, 3.3, 3.3], 9.2)
p('RSSI sebelum perubahan memakai sampel valid terakhir pada BSSID lama. Untuk perubahan pertama, sampel valid terakhir tercatat pada 16.21.06; dua pembacaan berikutnya, 16.21.07-16.21.08, tidak berisi BSSID. Pembacaan kosong dapat berkaitan dengan keadaan koneksi atau kegagalan pembacaan, sehingga tidak dianggap otomatis sebagai koneksi putus selama dua detik.')
p('Sebelum perubahan terakhir, RSSI koneksi turun menjadi -78 dBm, lalu tercatat -42 dBm pada radio baru. Ini menunjukkan klien berpindah setelah sinyal melemah, tetapi tidak cukup untuk menyatakan ada atau tidak ada sticky client. Log ping serentak dan lokasi pasti setiap perubahan tidak tersedia; packet loss, durasi gangguan, serta kesinambungan aplikasi saat roaming belum dapat dinilai.')

# Page 10
page('5 Rekomendasi perbaikan')
h('Prioritas 1 Memperbaiki cakupan menuju toilet', 2)
p('Area L8-011 menjadi prioritas utama karena RSSI WiFi-UB.x -73,0 dBm, ketiadaan AP cadangan, dan pengamatan koneksi eduroam yang lemah pada titik sama. Usulan awal adalah mengevaluasi pemindahan atau penambahan AP pada koridor sisi toilet, dengan posisi yang mengurangi penghalang menuju pintu dan area penggunaan. Lokasi pemasangan ditentukan melalui pemeriksaan lapangan; laporan ini tidak menetapkan jarak pemasangan dalam meter tanpa denah terukur.')
p('Setelah perubahan, ukur ulang L8-011 beserta koridor sekitarnya pada SSID WiFi-UB.x. Sasaran evaluasi adalah RSSI utama sekurang-kurangnya -67 dBm dan AP kedua sekurang-kurangnya -75 dBm, sambil memeriksa ulang latensi dan loss pada SSID yang sama. Perubahan daya pancar dilakukan bertahap, dengan verifikasi hasil agar tidak hanya memperkuat satu sisi cakupan.')
h('Prioritas 2 Memeriksa cakupan antar AP', 2)
p('L8-012 perlu diperiksa sebagai area dominasi satu radio WiFi-UB.x. Pastikan radio di sekitarnya memang memancarkan SSID target dan memiliki cakupan yang cukup pada jalur perpindahan. Ketiadaan AP kedua dalam tiga scan bukan alasan langsung untuk menurunkan daya AP-L8-02. Periksa pola cakupan, asosiasi klien, serta konfigurasi roaming sebelum menentukan penyesuaian.')
p('Uji ulang koneksi di area ini perlu memakai SSID WiFi-UB.x dan mencatat BSSID serta ping secara bersamaan. Hasil aktif yang terasosiasi ke eduroam pada kanal 48 belum dapat memastikan perilaku roaming pengguna WiFi-UB.x pada kanal 161.')
h('Prioritas 3 Mempertahankan kanal sambil mengevaluasi beban', 2)
p('Kombinasi kanal 40, 48, dan 161 dengan lebar 20 MHz dapat dipertahankan sebagai konfigurasi awal karena tidak ditemukan CCI tinggi atau utilisasi AP utama di atas 50% pada sesi ini. Belum ada bukti yang mengharuskan pelebaran kanal maupun perubahan kanal. Evaluasi pada jam penggunaan ramai dapat menentukan apakah bottleneck berasal dari kapasitas radio, layanan jaringan, atau cakupan.')
p('Jika AP baru ditambahkan menuju toilet, pilih kanal setelah memindai lingkungan di lokasi tersebut. Tujuannya menjaga cakupan dan menghindari overlap kanal dengan radio kuat di sekitarnya. Verifikasi kandidat pemetaan MAC AP melalui inventaris pengelola sebelum perubahan konfigurasi dilakukan.')

# Page 11
page('6 Keterbatasan dan kesimpulan')
h('6 1 Keterbatasan dan validasi', 2)
p('Jumlah titik dan waktu. Dua belas titik pada satu sesi sore belum mewakili seluruh ruang dan variasi beban, serta berada di bawah target 25-35 titik pada modul. Persentase hasil berlaku pada titik pengamatan.', 'Jumlah titik dan waktu.')
p('Denah dan interpolasi. Skala meter belum diverifikasi dan koordinat bersifat perkiraan. Heatmap tidak digunakan untuk menghitung luas cakupan atau jarak fisik presisi.', 'Denah dan interpolasi.')
p('Perangkat dan kalibrasi. Data L8-000 tidak tersedia. Pada L8-002, BSSID kanal 40 tercatat sekitar -50 dBm di Android dan -72,8 dBm rata-rata linear di laptop. Pengukuran yang tidak serentak ini bukan kalibrasi offset dan tidak dipakai untuk mengoreksi data.', 'Perangkat dan kalibrasi.')
p('Metrik yang tidak tersedia. Noise floor, SNR, dan throughput iperf3 tidak tersedia. Utilisasi berasal dari BSS Load saat scan, bukan pemantauan seluruh trafik secara terus-menerus.', 'Metrik yang tidak tersedia.')
p('SSID koneksi dan roaming. Data aktif dan roaming cocok dengan eduroam, sehingga tidak digeneralisasi ke WiFi-UB.x. Tanpa ping serentak, packet loss dan durasi gangguan handover belum dapat dinilai.', 'SSID koneksi dan roaming.')
p('Pemetaan AP. Pencocokan AP dengan BSSID bersifat kandidat; label MAC dan model belum terverifikasi. Asal radio dari lantai lain tidak dapat dipastikan tanpa data lintas lantai.', 'Pemetaan AP.')
h('6 2 Kesimpulan', 2)
p('Sebagian besar titik pengamatan lantai 8 memiliki cakupan WiFi-UB.x 5 GHz yang memadai: 11 dari 12 titik memenuhi RSSI >= -67 dBm. Status gabungan menunjukkan 10 BAIK, satu PERHATIAN, dan satu BURUK. Temuan utama adalah cakupan lemah di toilet L8-011 serta ketiadaan AP cadangan pada L8-011 dan L8-012. Prioritas tindakan adalah memperbaiki cakupan toilet dan mengevaluasi cakupan antar-AP, sedangkan perubahan kanal belum didukung oleh temuan CCI maupun utilisasi pada sesi ini.')
h('Sumber dan berkas pendukung', 2)
for text in [
    'Modul Hands-on Pemetaan Kualitas Sinyal Wi-Fi 5 GHz Gedung FILKOM, DRAFT v0.4, terutama bagian 4 dan 8-11.',
    'Data pengukuran Kelompok 4: passive_L8.csv, active_L8.csv, roaming_L8.csv, tangkapan WiFiAnalyzer, dan catatan ping mentah.',
    'Keluaran analyze.py: L8_titik.csv, L8_ap.csv, L8_area_lemah.csv, dan L8_ringkasan.txt. Gambar laporan dibuat ulang dari metrik ini dengan proporsi denah asli.',
    'Dokumentasi lapangan: sketsa_L8.jpg dan foto AP-L8-01, AP-L8-02, serta AP-L8-03.',
]:
    par = doc.add_paragraph(text)
    par.paragraph_format.space_after = Pt(4)
    for run in par.runs:
        run.font.size = Pt(9)

# Appendix A
page('Lampiran A Dokumentasi AP L8 01')
p('AP-L8-01 berada pada blok tengah bawah sesuai penomoran denah. Foto memperlihatkan perangkat pada plafon dengan tulisan Ruckus. Model dan MAC perangkat tidak terbaca pada dokumentasi ini.')
figure(photo_one, 'Gambar A1 Dokumentasi AP-L8-01 yang disampaikan pada 5 Oktober 2026', height=19.0)

# Appendix B
page('Lampiran B Dokumentasi AP L8 02 dan AP L8 03')
p('AP-L8-02 ditandai pada koridor kanan, sedangkan AP-L8-03 berada di bagian tengah atas denah. Tangkapan layar berikut memuat foto dan kode AP yang diberikan pada dokumentasi lapangan.')
figure(photo_other, 'Gambar B1 Foto AP-L8-02 pada bagian atas dan AP-L8-03 pada bagian bawah', height=18.8)
p('Dokumentasi ini menunjukkan keberadaan perangkat sesuai kode pengamatan. Label MAC dan model perangkat belum dapat digunakan untuk mengonfirmasi kandidat pemetaan BSSID.')

doc.core_properties.title = 'Laporan Pemetaan Kualitas Sinyal WiFi 5 GHz Lantai 8 FILKOM UB'
doc.core_properties.subject = 'Survei 28 September 2026 Kelompok 4 Jaringan Nirkabel'
doc.core_properties.author = 'Kelompok 4'
doc.core_properties.keywords = 'WiFi-UB.x, lantai 8, RSSI, access point, Jaringan Nirkabel'
doc.core_properties.comments = ''
path = OUT / 'Laporan_Survei_WiFi_L8_Kelompok4.docx'
doc.save(path)

sources = [BASE / 'data/passive_L8.csv', BASE / 'data/active_L8.csv', BASE / 'raw/roaming_L8.csv', BASE / 'output/L8_titik.csv', BASE / 'output/L8_ap.csv', photo_one, photo_other]
validation = {
    'passive_rows': len(passive),
    'passive_5ghz_rows': sum(r['band'] == '5' for r in passive),
    'target_5ghz_rows': len(target),
    'points': len(derived),
    'active_points': len(active),
    'roaming_samples': len(roam),
    'roaming_transitions': sum(r['pindah'] == 'YA' for r in roam),
    'roaming_empty_bssid': sum(not r['bssid'] for r in roam),
    'active_ssid_from_passive_bssid_match': 'eduroam',
    'rssi_independently_recomputed': True,
    'sources_sha256': {str(p.relative_to(BASE)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},
    'output_docx': str(path),
}
(WORK / 'validation.json').write_text(json.dumps(validation, indent=2), encoding='utf-8')
print(path)
print(json.dumps({k: v for k, v in validation.items() if k not in ['sources_sha256', 'output_docx']}, ensure_ascii=False))
